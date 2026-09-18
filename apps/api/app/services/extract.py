"""Text extraction per file type. Returns (full_text, pages).

Images: OCR is attempted when pytesseract + PIL are installed; otherwise a
clear OCRUnavailable error is raised and the document is marked FAILED with
an actionable message (never silently skipped).
"""
from __future__ import annotations

import re
from pathlib import Path


class ExtractionError(RuntimeError):
    pass


class OCRUnavailable(ExtractionError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: Path, file_type: str) -> tuple[str, list[str]]:
    if file_type == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = [_clean(page.extract_text() or "") for page in reader.pages]
        full = _clean("\n\n".join(pages))
        if not full:
            raise ExtractionError("No extractable text in PDF (scanned image?). "
                                  "OCR is not enabled for this deployment.")
        return full, pages

    if file_type == "docx":
        import docx
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        full = _clean("\n".join(parts))
        if not full:
            raise ExtractionError("No extractable text in DOCX.")
        return full, []

    if file_type in {"txt", "md", "csv"}:
        full = _clean(Path(path).read_text(encoding="utf-8", errors="replace"))
        if not full:
            raise ExtractionError("File contains no text.")
        return full, []

    if file_type == "image":
        try:
            import pytesseract
            from PIL import Image
        except ImportError as e:
            raise OCRUnavailable(
                "OCR engine (pytesseract) is not installed in this environment. "
                "Install it to process images, or upload a text-based file."
            ) from e
        text = _clean(pytesseract.image_to_string(Image.open(str(path))))
        if not text:
            raise ExtractionError("OCR found no text in the image.")
        return text, []

    raise ExtractionError(f"Unsupported file type: {file_type}")
